"""
Document Parser Service - 文档解析服务
支持 PDF、Word、Markdown 等格式
"""
import re
from typing import Dict, Any, Optional, List
from pathlib import Path


class DocumentParser:
    """文档解析器"""

    @staticmethod
    def parse_document(content: bytes, file_type: str, filename: str = "") -> Dict[str, Any]:
        """
        解析文档内容

        Args:
            content: 文件二进制内容
            file_type: 文件类型 (pdf/docx/md/txt)
            filename: 文件名

        Returns:
            {
                "text": str,  # 提取的文本内容
                "chunks": List[Dict],  # 分块结果
                "metadata": Dict  # 元数据
            }
        """
        file_type = file_type.lower()

        if file_type == "pdf":
            return DocumentParser._parse_pdf(content, filename)
        elif file_type in ["docx", "doc"]:
            return DocumentParser._parse_word(content, filename)
        elif file_type in ["md", "markdown"]:
            return DocumentParser._parse_markdown(content, filename)
        elif file_type == "txt":
            return DocumentParser._parse_text(content, filename)
        else:
            # 未知格式，尝试作为文本解析
            return DocumentParser._parse_text(content, filename)

    @staticmethod
    def _parse_pdf(content: bytes, filename: str = "") -> Dict[str, Any]:
        """解析 PDF 文档"""
        try:
            # 尝试使用 PyPDF2
            from PyPDF2 import PdfReader
            from io import BytesIO

            reader = PdfReader(BytesIO(content))
            text_parts = []
            chunks = []

            for i, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                text_parts.append(page_text)

                if page_text.strip():
                    chunks.append({
                        "id": f"chunk_{i+1}",
                        "content": page_text[:2000],  # 限制单块大小
                        "page": i + 1,
                        "type": "page"
                    })

            full_text = "\n\n".join(text_parts)

            return {
                "text": full_text,
                "chunks": chunks,
                "metadata": {
                    "filename": filename,
                    "type": "pdf",
                    "page_count": len(reader.pages),
                    "char_count": len(full_text)
                }
            }
        except ImportError:
            # PyPDF2 未安装，返回简化处理
            return DocumentParser._fallback_parse(content, filename, "pdf")
        except Exception as e:
            print(f"Error parsing PDF: {e}")
            return DocumentParser._fallback_parse(content, filename, "pdf")

    @staticmethod
    def _parse_word(content: bytes, filename: str = "") -> Dict[str, Any]:
        """解析 Word 文档 (.docx)"""
        try:
            from docx import Document
            from io import BytesIO

            doc = Document(BytesIO(content))
            text_parts = []
            chunks = []

            # 按段落解析
            current_chunk = ""
            chunk_index = 0

            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    text_parts.append(text)

                    # 累积段落到当前 chunk
                    if len(current_chunk) + len(text) < 1500:
                        current_chunk += text + "\n"
                    else:
                        # 保存当前 chunk，开始新 chunk
                        if current_chunk:
                            chunk_index += 1
                            chunks.append({
                                "id": f"chunk_{chunk_index}",
                                "content": current_chunk.strip(),
                                "type": "paragraph"
                            })
                        current_chunk = text + "\n"

            # 保存最后一个 chunk
            if current_chunk:
                chunk_index += 1
                chunks.append({
                    "id": f"chunk_{chunk_index}",
                    "content": current_chunk.strip(),
                    "type": "paragraph"
                })

            # 也尝试从表格中提取内容
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        text_parts.append(row_text)

            full_text = "\n\n".join(text_parts)

            return {
                "text": full_text,
                "chunks": chunks,
                "metadata": {
                    "filename": filename,
                    "type": "docx",
                    "paragraph_count": len(doc.paragraphs),
                    "table_count": len(doc.tables),
                    "char_count": len(full_text)
                }
            }
        except ImportError:
            return DocumentParser._fallback_parse(content, filename, "docx")
        except Exception as e:
            print(f"Error parsing Word document: {e}")
            return DocumentParser._fallback_parse(content, filename, "docx")

    @staticmethod
    def _parse_markdown(content: bytes, filename: str = "") -> Dict[str, Any]:
        """解析 Markdown 文档"""
        try:
            text = content.decode('utf-8', errors='ignore')

            # 移除 Markdown 标记，提取纯文本
            # 移除代码块
            text_no_code = re.sub(r'```[\s\S]*?```', '[代码块]', text)
            text_no_code = re.sub(r'`([^`]+)`', r'\1', text_no_code)

            # 移除链接标记，保留文本
            text_no_code = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text_no_code)

            # 移除图片标记
            text_no_code = re.sub(r'!\[([^\]]*)\]\([^)]+\)', '[图片]', text_no_code)

            # 移除标题标记 #
            text_no_code = re.sub(r'^#{1,6}\s+', '', text_no_code, flags=re.MULTILINE)

            # 移除粗体、斜体标记
            text_no_code = re.sub(r'\*\*\*([^*]+)\*\*\*', r'\1', text_no_code)
            text_no_code = re.sub(r'\*\*([^*]+)\*\*', r'\1', text_no_code)
            text_no_code = re.sub(r'\*([^*]+)\*', r'\1', text_no_code)
            text_no_code = re.sub(r'___([^_]+)___', r'\1', text_no_code)
            text_no_code = re.sub(r'__([^_]+)__', r'\1', text_no_code)
            text_no_code = re.sub(r'_([^_]+)_', r'\1', text_no_code)

            # 按标题分块
            chunks = []
            sections = re.split(r'\n(?=#{1,6}\s)', text)

            for i, section in enumerate(sections):
                section = section.strip()
                if section:
                    # 提取标题
                    title_match = re.match(r'^#{1,6}\s+(.+)$', section, re.MULTILINE)
                    title = title_match.group(1) if title_match else f"Section {i+1}"

                    # 截取内容预览
                    content_preview = section[:1500] if len(section) > 1500 else section

                    chunks.append({
                        "id": f"chunk_{i+1}",
                        "content": content_preview,
                        "title": title,
                        "type": "section"
                    })

            return {
                "text": text_no_code.strip(),
                "chunks": chunks[:50],  # 限制 chunk 数量
                "metadata": {
                    "filename": filename,
                    "type": "markdown",
                    "section_count": len(chunks),
                    "char_count": len(text)
                }
            }
        except Exception as e:
            print(f"Error parsing Markdown: {e}")
            return DocumentParser._fallback_parse(content, filename, "markdown")

    @staticmethod
    def _parse_text(content: bytes, filename: str = "") -> Dict[str, Any]:
        """解析纯文本文件"""
        try:
            text = content.decode('utf-8', errors='ignore')

            # 简单分块：按段落或固定长度
            chunks = []
            paragraphs = text.split('\n\n')

            current_chunk = ""
            chunk_index = 0

            for para in paragraphs:
                para = para.strip()
                if not para:
                    continue

                if len(current_chunk) + len(para) < 1500:
                    current_chunk += para + "\n\n"
                else:
                    if current_chunk:
                        chunk_index += 1
                        chunks.append({
                            "id": f"chunk_{chunk_index}",
                            "content": current_chunk.strip(),
                            "type": "paragraph"
                        })
                    current_chunk = para + "\n\n"

            if current_chunk:
                chunk_index += 1
                chunks.append({
                    "id": f"chunk_{chunk_index}",
                    "content": current_chunk.strip(),
                    "type": "paragraph"
                })

            return {
                "text": text.strip(),
                "chunks": chunks,
                "metadata": {
                    "filename": filename,
                    "type": "text",
                    "char_count": len(text)
                }
            }
        except Exception as e:
            print(f"Error parsing text: {e}")
            return DocumentParser._fallback_parse(content, filename, "text")

    @staticmethod
    def _fallback_parse(content: bytes, filename: str, file_type: str) -> Dict[str, Any]:
        """降级解析方案"""
        try:
            # 尝试作为文本解码
            text = content.decode('utf-8', errors='ignore')

            # 如果内容太长，只取前 10000 字符
            if len(text) > 10000:
                text = text[:10000] + "\n\n[内容已截断...]"

            return {
                "text": text,
                "chunks": [{
                    "id": "chunk_1",
                    "content": text[:2000],
                    "type": "fallback"
                }],
                "metadata": {
                    "filename": filename,
                    "type": file_type,
                    "char_count": len(text),
                    "fallback": True
                }
            }
        except Exception:
            return {
                "text": "",
                "chunks": [],
                "metadata": {
                    "filename": filename,
                    "type": file_type,
                    "error": "Failed to parse document"
                }
            }


# 全局解析器实例
document_parser = DocumentParser()
